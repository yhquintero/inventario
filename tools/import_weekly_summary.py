#!/usr/bin/env python3
"""Importa RESUMEN POR SEMANA PINAR.xlsx al estado compartido del servidor.

Uso: python3 tools/import_weekly_summary.py [--year 2026]

El resumen se guarda en state.weekly como referencia de Excel (excelSummary).
No reemplaza los movimientos ni los cálculos operativos de la aplicación.
"""
import json, re, os, sys, argparse, datetime, sqlite3
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX_PATH = os.path.join(ROOT, "xlsx", "RESUMEN POR SEMANA PINAR.xlsx")
DB_PATH = os.path.join(ROOT, "data", "cuadre.db")

MONTHS = {
    "ENERO": 1, "FEBRERO": 2, "MARZO": 3, "ABRIL": 4, "MAYO": 5, "JUNIO": 6,
    "JULIO": 7, "AGOSTO": 8, "SEPTIEMBRE": 9, "SETIEMBRE": 9,
    "OCTUBRE": 10, "NOVIEMBRE": 11, "DICIEMBRE": 12,
}

def norm_name(value):
    return re.sub(r"\s+", " ", str(value or "")).strip().upper()

def num(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    if value is None:
        return 0.0
    s = str(value).replace(",", "").replace(" ", "")
    try:
        return float(s)
    except ValueError:
        return 0.0

def round2(value):
    return round(value + 1e-10, 2)

def week_range(start_date):
    """Devuelve dict con from/to (lunes-domingo) para la semana de start_date."""
    d = datetime.date.fromisoformat(start_date)
    # lunes de esa semana
    monday = d - datetime.timedelta(days=d.weekday())
    sunday = monday + datetime.timedelta(days=6)
    return {"from": monday.isoformat(), "to": sunday.isoformat()}

def parse_weekly_summary(year):
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb["PINAR"]

    # Detectar mes
    month_name = ""
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=False):
        for cell in row:
            if cell.value:
                m = re.match(r"^\s*MES\s*:\s*(.+?)\s*$", str(cell.value), re.IGNORECASE)
                if m:
                    month_name = norm_name(m.group(1))
                    break
        if month_name:
            break
    month = MONTHS.get(month_name)
    if not month:
        raise ValueError(f"No se reconoció el mes: {month_name}")

    # Buscar fila de cabecera con "SEMANA N"
    header_row = -1
    label_col = -1
    total_col = -1
    week_cols = []  # lista de (label, col_index)
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True), 1):
        normalized = [norm_name(c) for c in row]
        candidates = [(label, idx) for idx, label in enumerate(normalized) if re.match(r"^SEMANA\s+\d+\s*\(", label)]
        if candidates:
            header_row = i
            label_col = normalized.index("GASTOS") if "GASTOS" in normalized else 0
            if "TOTAL" in normalized:
                total_col = normalized.index("TOTAL")
            week_cols = candidates
            break

    if header_row < 0 or not week_cols:
        raise ValueError("No se encontró la fila de cabecera con columnas SEMANA")

    # Leer filas de datos por etiqueta
    data_rows = {}
    for i in range(header_row + 1, ws.max_row + 1):
        row = [c.value for c in ws[i]]
        label = norm_name(row[label_col]).replace("  ", " ")
        if label:
            data_rows[label] = row

    def cell_by_label(label, col):
        row = data_rows.get(norm_name(label).replace("  ", " "))
        if row and col < len(row):
            return num(row[col])
        return 0.0

    expense_labels = [
        ("DOMICILIO", "domicilio"),
        ("LIMPIEZA", "limpieza"),
        ("CUSTODIO", "custodio"),
        ("SALARIO", "salario"),
        ("COMISIONES", "comisiones"),
        ("OTROS", "otros"),
    ]

    weeks = []
    for label, col in week_cols:
        # Extraer días: SEMANA N (inicio-fin)
        m = re.search(r"\((\d{1,2})\s*[-–]\s*(\d{1,2})\)", label)
        if not m:
            continue
        start_day = int(m.group(1))
        end_day = int(m.group(2))

        start_month = month
        start_year = year
        if start_day > end_day:
            start_month -= 1
            if start_month < 1:
                start_month = 12
                start_year -= 1

        first_day = datetime.date(start_year, start_month, start_day)
        if first_day.month != start_month:
            continue
        range_dates = week_range(first_day.isoformat())

        expenses_cup = {}
        for src, key in expense_labels:
            expenses_cup[key] = round2(cell_by_label(src, col))

        total_expenses_cup = round2(cell_by_label("TOTAL  DE GASTOS CUP", col))
        total_expenses_usd = round2(cell_by_label("TOTAL  DE GASTOS USD", col))
        sales_usd = round2(cell_by_label("TOTAL DE VENTA USD", col))
        gross_profit_usd = round2(cell_by_label("UTILIDAD EN VENTA USD", col))
        invested_usd = round2(cell_by_label("DINERO DE INVERSION USD", col))
        net_profit_usd = round2(cell_by_label("UTILIDAD NETA USD", col))

        cup_usd = round2(total_expenses_cup / total_expenses_usd) if total_expenses_usd > 0 else 0

        weeks.append({
            "weekLabel": label,
            "month": month_name,
            "year": year,
            "from": range_dates["from"],
            "to": range_dates["to"],
            "sourceFrom": first_day.isoformat(),
            "sourceToDay": end_day,
            "cupUsd": cup_usd,
            "expensesCup": expenses_cup,
            "totalExpensesCup": total_expenses_cup,
            "totalExpensesUsd": total_expenses_usd,
            "salesUsd": sales_usd,
            "grossProfitUsd": gross_profit_usd,
            "investedUsd": invested_usd,
            "netProfitUsd": net_profit_usd,
        })

    # Totales del libro
    totals = None
    if total_col >= 0:
        totals = {
            "expensesCup": {key: round2(cell_by_label(src, total_col)) for src, key in expense_labels},
            "totalExpensesCup": round2(cell_by_label("TOTAL  DE GASTOS CUP", total_col)),
            "totalExpensesUsd": round2(cell_by_label("TOTAL  DE GASTOS USD", total_col)),
            "salesUsd": round2(cell_by_label("TOTAL DE VENTA USD", total_col)),
            "grossProfitUsd": round2(cell_by_label("UTILIDAD EN VENTA USD", total_col)),
            "investedUsd": round2(cell_by_label("DINERO DE INVERSION USD", total_col)),
            "netProfitUsd": round2(cell_by_label("UTILIDAD NETA USD", total_col)),
        }

    return {"month": month_name, "year": year, "weeks": weeks, "totals": totals}


def update_server_state(parsed, imported_at):
    """Actualiza state.weekly en la base de datos del servidor."""
    con = sqlite3.connect(DB_PATH)
    con.row_factory = sqlite3.Row
    try:
        # Leer estado actual
        row = con.execute("SELECT version, data FROM app_state WHERE id=1").fetchone()
        if not row:
            raise RuntimeError("No hay estado en el servidor. Inicialice desde la Web primero.")
        version = row["version"]
        state = json.loads(row["data"]) if row["data"] else {}

        # Asegurar sección weekly
        if "weekly" not in state or not isinstance(state["weekly"], dict):
            state["weekly"] = {}

        # Guardar cada semana como excelSummary
        for week in parsed["weeks"]:
            key = week["from"]
            current = state["weekly"].get(key, {})
            ref = {**week, "sourceFile": "RESUMEN POR SEMANA PINAR.xlsx", "importedAt": imported_at}
            state["weekly"][key] = {**current, "excelSummary": ref}

        # Guardar estado actualizado
        new_version = version + 1
        con.execute(
            "UPDATE app_state SET version=?, data=?, updated_at=?, updated_by=? WHERE id=1",
            (new_version, json.dumps(state, ensure_ascii=False, separators=(",", ":")), int(__import__("time").time()), "import_weekly_summary.py")
        )
        con.commit()
        print(f"OK Estado actualizado: version {version} -> {new_version}")
        print(f"OK {len(parsed['weeks'])} semanas guardadas en state.weekly")
        for w in parsed["weeks"]:
            print(f"  - {w['weekLabel']}: {w['from']} a {w['to']} - Ventas ${w['salesUsd']:,.2f} - Utilidad neta ${w['netProfitUsd']:,.2f}")
    finally:
        con.close()


def main():
    parser = argparse.ArgumentParser(description="Importar resumen semanal PINAR al servidor")
    parser.add_argument("--year", type=int, default=datetime.date.today().year, help="Año del resumen (el Excel no lo incluye)")
    args = parser.parse_args()

    print(f"Leyendo {XLSX_PATH}...")
    parsed = parse_weekly_summary(args.year)
    print(f"Mes: {parsed['month']} {parsed['year']} - {len(parsed['weeks'])} semanas detectadas")

    imported_at = int(__import__("time").time() * 1000)
    update_server_state(parsed, imported_at)

    print("\nListo! El resumen esta disponible en la Web (Informe semanal -> ver referencia Excel)")
    print("y en la App tras la proxima sincronizacion (la App aun no tiene UI para verlo).")


if __name__ == "__main__":
    main()